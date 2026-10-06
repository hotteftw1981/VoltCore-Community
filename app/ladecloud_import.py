import hashlib
import io
import re
from collections import defaultdict
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

REQUIRED_HEADERS = {
    "Startdatum (UTC)", "Dauer", "kWh", "Ladepunkt", "Anschluss",
    "EVSE ID", "Nutzungsart", "RFID Tag", "RFID Nutzer",
}


def _text(value):
    return str(value or "").strip()


def _number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw=str(value).strip().replace(".", "").replace(",", ".")
    try:
        return float(raw)
    except ValueError:
        return None


def _duration_seconds(value):
    if value in (None, ""):
        return 0
    if isinstance(value, (int, float)):
        # Excel duration values are fractions of a day.
        return max(0, int(round(float(value) * 86400)))
    raw=str(value).strip().lower()
    h=re.search(r"(\d+)\s*h",raw); m=re.search(r"(\d+)\s*m",raw); s=re.search(r"(\d+)\s*s",raw)
    if not any((h,m,s)):
        parts=raw.split(":")
        if len(parts) in (2,3) and all(x.strip().isdigit() for x in parts):
            nums=[int(x) for x in parts]
            if len(nums)==2: return nums[0]*60+nums[1]
            return nums[0]*3600+nums[1]*60+nums[2]
        return 0
    return (int(h.group(1)) if h else 0)*3600+(int(m.group(1)) if m else 0)*60+(int(s.group(1)) if s else 0)


def _start_utc(value):
    if isinstance(value, datetime):
        dt=value
    else:
        raw=_text(value)
        if not raw: return None
        dt=None
        for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                dt=datetime.strptime(raw,fmt); break
            except ValueError:
                pass
        if dt is None:
            try: dt=datetime.fromisoformat(raw.replace("Z","+00:00"))
            except ValueError: return None
    if dt.tzinfo is None: dt=dt.replace(tzinfo=timezone.utc)
    else: dt=dt.astimezone(timezone.utc)
    return dt


def _find_header(ws):
    for row_idx in range(1, min(ws.max_row, 30)+1):
        values=[_text(c.value) for c in ws[row_idx]]
        if "Startdatum (UTC)" in values and "RFID Tag" in values:
            return row_idx, values
    raise ValueError("Die lade.cloud-Spaltenüberschrift wurde nicht gefunden.")


def parse_ladecloud_xlsx(data: bytes):
    if not data or len(data) > 20 * 1024 * 1024:
        raise ValueError("Die XLSX-Datei ist leer oder größer als 20 MB.")
    try:
        wb=load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Die Datei konnte nicht als XLSX gelesen werden.") from exc
    ws=wb.active
    header_row, headers=_find_header(ws)
    missing=sorted(REQUIRED_HEADERS-set(headers))
    if missing:
        raise ValueError("Im lade.cloud-Export fehlen Spalten: "+", ".join(missing))
    idx={name:i for i,name in enumerate(headers) if name}
    rows=[]; users=defaultdict(lambda:{"sessions":0,"energy_kwh":0.0,"source_names":set()}); cps=defaultdict(lambda:{"sessions":0,"energy_kwh":0.0,"evse_ids":set(),"connectors":set()})
    warnings=[]
    for row_no, row in enumerate(ws.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
        start=_start_utc(row[idx["Startdatum (UTC)"]] if idx["Startdatum (UTC)"] < len(row) else None)
        if start is None: continue
        usage=_text(row[idx["Nutzungsart"]] if idx["Nutzungsart"] < len(row) else None)
        if usage and usage.casefold() != "rfid":
            continue
        tag=_text(row[idx["RFID Tag"]] if idx["RFID Tag"] < len(row) else None)
        source_user=_text(row[idx["RFID Nutzer"]] if idx["RFID Nutzer"] < len(row) else None)
        source_cp=_text(row[idx["Ladepunkt"]] if idx["Ladepunkt"] < len(row) else None)
        if not tag or not source_cp:
            warnings.append(f"Zeile {row_no}: RFID oder Ladepunkt fehlt; Zeile übersprungen.")
            continue
        energy=_number(row[idx["kWh"]] if idx["kWh"] < len(row) else None) or 0.0
        duration=_duration_seconds(row[idx["Dauer"]] if idx["Dauer"] < len(row) else None)
        connector=int(_number(row[idx["Anschluss"]] if idx["Anschluss"] < len(row) else None) or 1)
        evse=_text(row[idx["EVSE ID"]] if idx["EVSE ID"] < len(row) else None)
        price_eur=_number(row[idx["RFID Preis €/kWh"]] if "RFID Preis €/kWh" in idx and idx["RFID Preis €/kWh"] < len(row) else None)
        end=start+timedelta(seconds=duration)
        started_at=start.isoformat(); ended_at=end.isoformat()
        raw_key="|".join([started_at,tag,source_cp,str(connector),evse,f"{energy:.6f}",str(duration)])
        import_key=hashlib.sha256(raw_key.encode("utf-8")).hexdigest()
        item={
            "row_no":row_no,"started_at":started_at,"ended_at":ended_at,"duration_seconds":duration,
            "energy_kwh":round(float(energy),6),"source_charge_point":source_cp,"connector_id":connector,"evse_id":evse,
            "rfid_tag":tag,"source_user_name":source_user,"price_cents_per_kwh":None if price_eur is None else int(round(price_eur*100)),
            "cost_cents":0,"import_key":import_key,
        }
        rows.append(item)
        u=users[tag]; u["sessions"]+=1; u["energy_kwh"]+=energy; u["source_names"].add(source_user or tag)
        c=cps[source_cp]; c["sessions"]+=1; c["energy_kwh"]+=energy; c["evse_ids"].add(evse); c["connectors"].add(connector)
    if not rows:
        raise ValueError("Der Export enthält keine importierbaren RFID-Ladevorgänge.")
    user_list=[]
    for tag,u in sorted(users.items(), key=lambda x: sorted(x[1]["source_names"])[0].casefold()):
        user_list.append({"rfid_tag":tag,"source_user_name":sorted(u["source_names"])[0],"sessions":u["sessions"],"energy_kwh":round(u["energy_kwh"],3)})
    cp_list=[]
    for name,c in sorted(cps.items()):
        cp_list.append({"source_charge_point":name,"sessions":c["sessions"],"energy_kwh":round(c["energy_kwh"],3),"evse_ids":sorted(x for x in c["evse_ids"] if x),"connectors":sorted(c["connectors"])})
    return {
        "rows":rows,"users":user_list,"charge_points":cp_list,"sessions":len(rows),"energy_kwh":round(sum(x["energy_kwh"] for x in rows),3),
        "first_at":min(x["started_at"] for x in rows),"last_at":max(x["started_at"] for x in rows),"warnings":warnings[:25],
    }
